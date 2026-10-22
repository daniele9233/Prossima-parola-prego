"""Impagina prologo + capitolo con riquadri e blocchi di codice."""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Palatino Linotype"
MONO = "Consolas"
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
QUOTE = RGBColor(0x2E, 0x2E, 0x2E)
RULE = "8A8A8A"

TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*)")

BOXES = {
    "cofano": ("Sotto il cofano", "F4F2EC", "2F2F2F"),
    "prova": ("Provalo tu", "FBF4E4", "B08D3F"),
    "curiosita": ("Curiosità", "F7F6F2", "C4BEAF"),
    "treghe": ("In tre righe", "EFEFEC", "2F2F2F"),
}


def typographic(text: str) -> str:
    text = text.replace("'", "\u2019")
    text = re.sub(r"(\u00ab)\s+", "\u00ab", text)
    text = re.sub(r"\s+(\u00bb)", "\u00bb", text)
    return text


def letterspace(run, hundredths: int) -> None:
    rpr = run._element.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(hundredths))
    rpr.append(spacing)


def shade(par, color: str) -> None:
    ppr = par._p.get_or_add_pPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), color)
    ppr.insert(0, el)


def border(par, color: str, sides=("left",), sz=18, space=8) -> None:
    ppr = par._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for side in sides:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color)
        pbdr.append(el)
    ppr.insert(0, pbdr)


def add_runs(par, text: str, size=None) -> None:
    for part in TOKEN.split(typographic(text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = par.add_run(part[1:-1])
            run.italic = True
        else:
            run = par.add_run(part)
        if size:
            run.font.size = Pt(size)


def page_number(par) -> None:
    run = par.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    run.font.name = FONT
    run.font.size = Pt(9)
    run.font.color.rgb = GREY


def base_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(15.24)
    section.page_height = Cm(22.86)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(1.8)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.42
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(0.5)
    pf.widow_control = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_number(footer)
    return doc


def flush_code(box) -> None:
    if box["code"]:
        box["items"].append(("code", list(box["code"])))
        box["code"] = []


def parse(md_path: Path):
    raw = md_path.read_text(encoding="utf-8").splitlines()
    sections = []
    current = None
    box = None

    for line in raw:
        line = line.rstrip()

        if box is not None:
            if line.strip() == ":::":
                flush_code(box)
                sections[-1]["blocks"].append(
                    ("box", box["kind"], box["title"], box["items"])
                )
                box = None
                continue
            if line.startswith("    "):
                box["code"].append(line[4:])
                continue
            flush_code(box)
            if not line.strip():
                continue
            box["items"].append(("t", line.strip()))
            continue

        if line.startswith(":::") and line.strip() != ":::":
            head = line[3:].strip()
            kind, _, title = head.partition(" ")
            box = {"kind": kind, "title": title.strip(), "items": [], "code": []}
            continue

        if line.startswith("# "):
            current = {"title": line[2:].strip(), "blocks": []}
            sections.append(current)
            continue
        if line.startswith("## "):
            current["blocks"].append(("h2", line[3:].strip()))
            continue
        if not line.strip():
            continue
        if line.strip() == "---":
            continue
        current["blocks"].append(("p", line.strip()))

    return sections


def title_block(doc, label: str, title: str, first: bool) -> None:
    if not first:
        doc.add_page_break()

    stamp = doc.add_paragraph()
    stamp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    stamp.paragraph_format.first_line_indent = Cm(0)
    stamp.paragraph_format.space_before = Pt(0 if first else 46)
    stamp.paragraph_format.space_after = Pt(16)
    run = stamp.add_run(label.upper())
    run.font.size = Pt(9)
    run.font.color.rgb = GREY
    letterspace(run, 60)

    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    head.paragraph_format.first_line_indent = Cm(0)
    head.paragraph_format.line_spacing = 1.1
    head.paragraph_format.space_after = Pt(10)
    head.paragraph_format.keep_with_next = True
    run = head.add_run(typographic(title))
    run.font.size = Pt(19)

    orn = doc.add_paragraph()
    orn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    orn.paragraph_format.first_line_indent = Cm(0)
    orn.paragraph_format.space_after = Pt(24)
    run = orn.add_run("*  *  *")
    run.font.size = Pt(10)
    run.font.color.rgb = GREY


def render_box(doc, kind: str, title: str, items) -> None:
    if kind == "pensa":
        render_pensa(doc, title, items)
        return

    label, fill, edge = BOXES.get(kind, ("", "F4F2EC", "2F2F2F"))
    low = title.lower()
    if low.startswith(label.lower()):
        title = title[len(label):].lstrip(" —-·:")

    def new_par(first: bool, last: bool):
        par = doc.add_paragraph()
        par.paragraph_format.first_line_indent = Cm(0)
        par.paragraph_format.left_indent = Cm(0.45)
        par.paragraph_format.right_indent = Cm(0.25)
        par.paragraph_format.line_spacing = 1.25
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        par.paragraph_format.space_before = Pt(8 if first else 0)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.keep_together = True
        par.paragraph_format.keep_with_next = not last
        if last:
            par.paragraph_format.space_after = Pt(12)
        shade(par, fill)
        border(par, edge)
        return par

    # etichetta
    lab_par = new_par(True, False)
    lab = lab_par.add_run(label.upper())
    lab.bold = True
    lab.font.size = Pt(8)
    lab.font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)
    letterspace(lab, 30)

    # titolo del riquadro
    if title:
        tip = new_par(False, False)
        tip.paragraph_format.space_before = Pt(3)
        tip.paragraph_format.space_after = Pt(4)
        tip.paragraph_format.keep_with_next = True
        run = tip.add_run(typographic(title))
        run.bold = True
        run.font.size = Pt(10)

    for item in items:
        if item[0] == "t":
            par = new_par(False, False)
            add_runs(par, item[1], size=9.5)
        elif item[0] == "code":
            par = new_par(False, False)
            par.paragraph_format.left_indent = Cm(0.75)
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.space_before = Pt(6)
            lines = item[1]
            for idx, code_line in enumerate(lines):
                run = par.add_run(code_line)
                run.font.name = MONO
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                if idx < len(lines) - 1:
                    run.add_break()

    # chiusura del riquadro
    par = new_par(False, True)
    run = par.add_run("")
    run.font.size = Pt(2)


def render_pensa(doc, title: str, items) -> None:
    label = "Per chi ama pensare"
    if title.lower().startswith(label.lower()):
        title = title[len(label):].lstrip(" —-·:")

    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    head.paragraph_format.first_line_indent = Cm(0)
    head.paragraph_format.space_before = Pt(16)
    head.paragraph_format.space_after = Pt(2)
    head.paragraph_format.keep_with_next = True
    border(head, RULE, sides=("top",), sz=6, space=10)
    run = head.add_run("PER CHI AMA PENSARE")
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    letterspace(run, 40)

    if title and title != "Per chi ama pensare":
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.paragraph_format.first_line_indent = Cm(0)
        sub.paragraph_format.space_after = Pt(8)
        sub.paragraph_format.keep_with_next = True
        run = sub.add_run(typographic(title))
        run.italic = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = GREY

    last = None
    for item in items:
        if item[0] != "t":
            continue
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.first_line_indent = Cm(0)
        par.paragraph_format.left_indent = Cm(0.6)
        par.paragraph_format.right_indent = Cm(0.6)
        par.paragraph_format.line_spacing = 1.3
        par.paragraph_format.space_after = Pt(7)
        par.paragraph_format.keep_together = True
        add_runs(par, item[1], size=10)
        for run in par.runs:
            run.italic = True
            run.font.color.rgb = QUOTE
        last = par

    if last is not None:
        last.paragraph_format.space_after = Pt(18)
        border(last, RULE, sides=("bottom",), sz=6, space=10)


def build(md_path: Path, docx_path: Path, stamp: str) -> None:
    sections = parse(md_path)
    doc = base_document()

    first_section = True
    for sec in sections:
        label, _, title = sec["title"].partition(" — ")
        if not title:
            label, title = "", sec["title"]
        title_block(doc, label or sec["title"], title or sec["title"], first_section)
        first_section = False

        first_para = True
        for block in sec["blocks"]:
            if block[0] == "h2":
                par = doc.add_paragraph()
                par.paragraph_format.first_line_indent = Cm(0)
                par.paragraph_format.space_before = Pt(17)
                par.paragraph_format.space_after = Pt(7)
                par.paragraph_format.keep_with_next = True
                par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = par.add_run(typographic(block[1]))
                run.bold = True
                run.font.size = Pt(11.5)
                first_para = True
                continue

            if block[0] == "box":
                render_box(doc, block[1], block[2], block[3])
                first_para = True
                continue

            text = block[1].strip()
            if text.startswith("**") and text.endswith("**") and text.count("**") == 2:
                par = doc.add_paragraph()
                par.paragraph_format.first_line_indent = Cm(0)
                par.paragraph_format.left_indent = Cm(0.7)
                par.paragraph_format.right_indent = Cm(0.7)
                par.paragraph_format.space_before = Pt(15)
                par.paragraph_format.space_after = Pt(15)
                par.paragraph_format.line_spacing = 1.3
                par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = par.add_run(typographic(text[2:-2]))
                run.italic = True
                run.font.size = Pt(11.5)
                run.font.color.rgb = QUOTE
                first_para = True
                continue

            par = doc.add_paragraph()
            if first_para:
                par.paragraph_format.first_line_indent = Cm(0)
                first_para = False
            add_runs(par, text)

    props = doc.core_properties
    props.title = sections[0]["title"]
    props.subject = stamp
    props.author = ""
    doc.save(docx_path)


if __name__ == "__main__":
    build(Path(sys.argv[1]), Path(sys.argv[2]), sys.argv[3])
